import openpyxl
import sys
import io

# Redirect stdout to handle UTF-8
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('ARCHER_Enhanced_Features_FINAL.xlsx')

if 'Cross-Agent Guidelines' in wb.sheetnames:
    sheet = wb['Cross-Agent Guidelines']
    print('=== Cross-Agent Guidelines ===')
    for row in sheet.iter_rows(values_only=True):
        if row and any(cell for cell in row if cell):
            print(row)

if 'Dev Agents Plan' in wb.sheetnames:
    sheet = wb['Dev Agents Plan']
    print('\n=== Dev Agents Plan ===')
    for row in sheet.iter_rows(values_only=True):
        if row and any(cell for cell in row if cell):
            print(row)